from openpyxl import Workbook, load_workbook
from datetime import datetime
import tkinter as tk
from tkinter import messagebox
import os

window = tk.Tk()
window.title("Indy's Data Collector")
window.geometry("700x500")
window.iconbitmap("Icon.ico")

patient_name = ""

# ==========================
# Excel File Setup
# ==========================

FILE_NAME = "patient_data.xlsx"


def save_data():
    name = patient_name
    date = date_entry.get()
    systolic = systolic_entry.get()
    diastolic = diastolic_entry.get()
    heart_rate = hr_entry.get()

    if not name or not date or not systolic or not diastolic or not heart_rate:
        messagebox.showerror(
            "Missing Information",
            "Please fill in all fields."
        )
        return

    if os.path.exists(FILE_NAME):
        workbook = load_workbook(FILE_NAME)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active

        sheet.title = "Patient Data"

        sheet.cell(row=1, column=1).value = "Patient Name"
        sheet.cell(row=1, column=2).value = "Date"
        sheet.cell(row=1, column=3).value = "Systolic BP"
        sheet.cell(row=1, column=4).value = "Diastolic BP"
        sheet.cell(row=1, column=5).value = "Heart Rate"

    next_row = sheet.max_row + 1

    sheet.cell(row=next_row, column=1).value = name
    sheet.cell(row=next_row, column=2).value = date
    sheet.cell(row=next_row, column=3).value = systolic
    sheet.cell(row=next_row, column=4).value = diastolic
    sheet.cell(row=next_row, column=5).value = heart_rate

    workbook.save(FILE_NAME)

    messagebox.showinfo(
        "Success",
        "Patient data saved successfully!"
    )

    date_entry.delete(0, tk.END)
    systolic_entry.delete(0, tk.END)
    diastolic_entry.delete(0, tk.END)
    hr_entry.delete(0, tk.END)


# ==========================
# Screen Switching
# ==========================

def show_data_screen():
    global patient_name

    patient_name = patient_name_entry.get().strip()

    if not patient_name:
        messagebox.showerror(
            "Missing Name",
            "Please enter the patient's name."
        )
        return

    patient_name_label.config(
        text=f"Patient: {patient_name}"
    )

    name_screen.pack_forget()
    data_screen.pack(fill="both", expand=True)


def return_to_name_screen():
    data_screen.pack_forget()
    name_screen.pack(fill="both", expand=True)


# ==========================
# Name Screen
# ==========================

name_screen = tk.Frame(window)

name_title = tk.Label(
    name_screen,
    text="Health-E",
    font=("Arial", 20, "bold")
)
name_title.pack(pady=30)

name_instruction = tk.Label(
    name_screen,
    text="Please enter the patient's name:"
)
name_instruction.pack(pady=10)

patient_name_entry = tk.Entry(
    name_screen,
    width=35,
    font=("Arial", 12)
)
patient_name_entry.pack(pady=10)

continue_button = tk.Button(
    name_screen,
    text="Continue",
    command=show_data_screen,
    width=15
)
continue_button.pack(pady=20)

name_screen.pack(fill="both", expand=True)


# ==========================
# Data Screen
# ==========================

data_screen = tk.Frame(window)

patient_name_label = tk.Label(
    data_screen,
    text="Patient:",
    font=("Arial", 16, "bold")
)
patient_name_label.pack(pady=20)

data_title = tk.Label(
    data_screen,
    text="Patient Measurements",
    font=("Arial", 18)
)
data_title.pack(pady=10)

# Date

date_frame = tk.Frame(data_screen)
date_frame.pack(pady=10)

date_label = tk.Label(
    date_frame,
    text="Date:",
    width=15,
    anchor="e"
)
date_label.pack(side="left", padx=5)

date_entry = tk.Entry(
    date_frame,
    width=25
)
date_entry.pack(side="left", padx=5)

date_entry.insert(
    0,
    datetime.now().strftime("%d/%m/%Y")
)

# Blood Pressure

bp_frame = tk.Frame(data_screen)
bp_frame.pack(pady=10)

bp_label = tk.Label(
    bp_frame,
    text="Blood Pressure:",
    width=15,
    anchor="e"
)
bp_label.pack(side="left", padx=5)

systolic_entry = tk.Entry(
    bp_frame,
    width=6
)
systolic_entry.pack(side="left", padx=2)

slash_label = tk.Label(
    bp_frame,
    text="/",
    font=("Arial", 12, "bold")
)
slash_label.pack(side="left", padx=2)

diastolic_entry = tk.Entry(
    bp_frame,
    width=6
)
diastolic_entry.pack(side="left", padx=2)

bp_unit_label = tk.Label(
    bp_frame,
    text="mmHg"
)
bp_unit_label.pack(side="left", padx=5)

# Heart Rate

hr_frame = tk.Frame(data_screen)
hr_frame.pack(pady=10)

hr_label = tk.Label(
    hr_frame,
    text="Heart Rate:",
    width=15,
    anchor="e"
)
hr_label.pack(side="left", padx=5)

hr_entry = tk.Entry(
    hr_frame,
    width=10
)
hr_entry.pack(side="left", padx=5)

hr_unit_label = tk.Label(
    hr_frame,
    text="BPM"
)
hr_unit_label.pack(side="left", padx=5)

# Buttons

button_frame = tk.Frame(data_screen)
button_frame.pack(pady=30)

save_button = tk.Button(
    button_frame,
    text="Save Data",
    command=save_data,
    width=15
)
save_button.pack(side="left", padx=10)

back_button = tk.Button(
    button_frame,
    text="Change Patient",
    command=return_to_name_screen,
    width=15
)
back_button.pack(side="left", padx=10)

window.mainloop()