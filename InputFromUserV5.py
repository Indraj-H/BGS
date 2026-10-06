from openpyxl import Workbook
import tkinter as tk

window = tk.Tk()
window.title("Indy's Data Collector")
window.geometry("250x200")

# Title
title = tk.Label(window, text="Welcome to Indy's Data Collector")
title.pack()

# Name inputs
label1 = tk.Label(window, text="Please insert person 1 fullname:")
label1.pack()

name1 = tk.Entry(window)
name1.pack()

label2 = tk.Label(window, text="Please insert person 2 fullname:")
label2.pack()

name2 = tk.Entry(window)
name2.pack()

label3 = tk.Label(window, text="Please insert person 3 fullname:")
label3.pack()

name3 = tk.Entry(window)
name3.pack()

# Function to collect the data
def collect_data():

    # Get the names AFTER the user has typed them
    person1 = name1.get()
    person2 = name2.get()
    person3 = name3.get()

    # Excel Workbook Initialization
    workbook = Workbook()
    sheet = workbook.active

    # Put names into Excel
    sheet.cell(row=1, column=1).value = "Name"
    sheet.cell(row=2, column=1).value = person1
    sheet.cell(row=3, column=1).value = person2
    sheet.cell(row=4, column=1).value = person3

    # Save Excel file
    workbook.save("names.xlsx")

    # Display success message in the window
    success = tk.Label(window, text="Data Collected Successfully!", fg="green")
    success.pack()

# Button
button = tk.Button(window, text="Collect Data", command=collect_data)
button.pack()

# Keep window open
window.mainloop()