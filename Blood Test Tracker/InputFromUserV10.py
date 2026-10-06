from openpyxl import Workbook, load_workbook
from datetime import datetime
import tkinter as tk
from tkinter import messagebox
import os
import re


window = tk.Tk()
window.title("Indy's Data Collector")
window.geometry("700x500")
window.iconbitmap("Icon.ico")

patient_name = ""
file_name = ""

# ==========================
# File Location
# ==========================

FILE_PATH = r"C:\Users\indra\OneDrive\HSH\Health-E"

os.makedirs(FILE_PATH, exist_ok=True)


# ==========================
# Patient Names
# ==========================

def get_existing_names():
    names = []

    for file in os.listdir(FILE_PATH):
        if file.endswith(".xlsx"):
            name = os.path.splitext(file)[0]
            names.append(name)

    return sorted(names)


existing_names = get_existing_names()


def create_file_name(name):
    safe_name = re.sub(r'[\\/*?:"<>|]', '', name)
    safe_name = safe_name.strip()

    if not safe_name:
        safe_name = "Unknown Patient"

    return f"{safe_name}.xlsx"


def update_dropdown(event=None):
    typed_name = patient_name_entry.get().lower()

    matching_names = [
        name for name in existing_names
        if typed_name in name.lower()
    ]

    suggestion_list.delete(0, tk.END)

    if typed_name and matching_names:
        for name in matching_names:
            suggestion_list.insert(tk.END, name)

        suggestion_list.pack(pady=5)
    else:
        suggestion_list.pack_forget()


def select_name(event=None):
    selected = suggestion_list.curselection()

    if selected:
        selected_name = suggestion_list.get(selected[0])

        patient_name_entry.delete(0, tk.END)
        patient_name_entry.insert(0, selected_name)

        suggestion_list.pack_forget()


# ==========================
# Risk Warning
# ==========================

def check_risk(event=None):

    systolic = systolic_entry.get().strip()
    diastolic = diastolic_entry.get().strip()
    heart_rate = hr_entry.get().strip()

    # Hide warning if fields are empty
    if not systolic or not diastolic or not heart_rate:
        risk_warning.pack_forget()
        return

    try:
        systolic_value = int(systolic)
        diastolic_value = int(diastolic)
        heart_rate_value = int(heart_rate)

    except ValueError:
        risk_warning.pack_forget()
        return

    # ==========================
    # Risk Limits
    # ==========================

    patient_at_risk = (
        systolic_value > 150
        or systolic_value < 100
        or diastolic_value > 90
        or diastolic_value < 70
        or heart_rate_value > 120
        or heart_rate_value < 50
    )

    if patient_at_risk:

        risk_warning.pack(
            pady=10,
            padx=20,
            fill="x"
        )

    else:

        risk_warning.pack_forget()


# ==========================
# Save Data
# ==========================

def save_data():

    date = date_entry.get().strip()
    systolic = systolic_entry.get().strip()
    diastolic = diastolic_entry.get().strip()
    heart_rate = hr_entry.get().strip()

    if not date or not systolic or not diastolic or not heart_rate:

        messagebox.showerror(
            "Missing Information",
            "Please fill in all fields."
        )

        return

    # ==========================
    # Check Date
    # ==========================

    try:

        datetime.strptime(
            date,
            "%d/%m/%Y"
        )

    except ValueError:

        messagebox.showerror(
            "Invalid Date",
            "Please enter the date in DD/MM/YYYY format."
        )

        return

    # ==========================
    # Check Numbers
    # ==========================

    try:

        systolic_value = int(systolic)
        diastolic_value = int(diastolic)
        heart_rate_value = int(heart_rate)

    except ValueError:

        messagebox.showerror(
            "Invalid Information",
            "Blood pressure and heart rate must be numbers."
        )

        return

    # ==========================
    # File Path
    # ==========================

    full_file_path = os.path.join(
        FILE_PATH,
        file_name
    )

    # ==========================
    # Open Existing File
    # ==========================

    if os.path.exists(full_file_path):

        workbook = load_workbook(
            full_file_path
        )

        sheet = workbook.active

    else:

        workbook = Workbook()

        sheet = workbook.active

        sheet.title = "Patient Data"

        sheet.cell(
            row=1,
            column=1
        ).value = "Date"

        sheet.cell(
            row=1,
            column=2
        ).value = "Systolic BP"

        sheet.cell(
            row=1,
            column=3
        ).value = "Diastolic BP"

        sheet.cell(
            row=1,
            column=4
        ).value = "Heart Rate"

    # ==========================
    # Add New Data
    # ==========================

    next_row = sheet.max_row + 1

    sheet.cell(
        row=next_row,
        column=1
    ).value = date

    sheet.cell(
        row=next_row,
        column=2
    ).value = systolic_value

    sheet.cell(
        row=next_row,
        column=3
    ).value = diastolic_value

    sheet.cell(
        row=next_row,
        column=4
    ).value = heart_rate_value

    # ==========================
    # Sort Dates
    # Oldest to Newest
    # ==========================

    data = list(
        sheet.iter_rows(
            min_row=2,
            values_only=True
        )
    )

    try:

        data.sort(
            key=lambda row: datetime.strptime(
                str(row[0]),
                "%d/%m/%Y"
            )
        )

    except ValueError:

        messagebox.showerror(
            "Date Error",
            "One of the dates in the file is invalid."
        )

        workbook.close()

        return

    # ==========================
    # Rewrite Sorted Data
    # ==========================

    for row_number, row_data in enumerate(
        data,
        start=2
    ):

        for column_number, value in enumerate(
            row_data,
            start=1
        ):

            sheet.cell(
                row=row_number,
                column=column_number
            ).value = value

    # ==========================
    # Save Workbook
    # ==========================

    try:

        workbook.save(
            full_file_path
        )

        workbook.close()

    except PermissionError:

        workbook.close()

        messagebox.showerror(
            "File Error",
            "The Excel file is currently open.\n\n"
            "Please close the file in Excel and try again."
        )

        return

    # ==========================
    # Success Message
    # ==========================

    messagebox.showinfo(
        "Success",
        f"Data saved to {file_name}!"
    )

    # ==========================
    # Clear Measurement Fields
    # ==========================

    date_entry.delete(
        0,
        tk.END
    )

    systolic_entry.delete(
        0,
        tk.END
    )

    diastolic_entry.delete(
        0,
        tk.END
    )

    hr_entry.delete(
        0,
        tk.END
    )

    date_entry.insert(
        0,
        datetime.now().strftime(
            "%d/%m/%Y"
        )
    )

    risk_warning.pack_forget()


# ==========================
# Screen Switching
# ==========================

def show_data_screen():

    global patient_name
    global file_name
    global existing_names

    patient_name = patient_name_entry.get().strip()

    if not patient_name:

        messagebox.showerror(
            "Missing Name",
            "Please enter the patient's name."
        )

        return

    file_name = create_file_name(
        patient_name
    )

    if patient_name not in existing_names:

        existing_names.append(
            patient_name
        )

        existing_names.sort()

    patient_name_label.config(
        text=f"Patient: {patient_name}"
    )

    name_screen.pack_forget()

    data_screen.pack(
        fill="both",
        expand=True
    )


def return_to_name_screen():

    data_screen.pack_forget()

    patient_name_entry.delete(
        0,
        tk.END
    )

    suggestion_list.pack_forget()

    name_screen.pack(
        fill="both",
        expand=True
    )


# ==========================
# Name Screen
# ==========================

name_screen = tk.Frame(
    window
)

name_title = tk.Label(
    name_screen,
    text="Health-E",
    font=("Arial", 20, "bold")
)

name_title.pack(
    pady=30
)

name_instruction = tk.Label(
    name_screen,
    text="Please enter the patient's name:"
)

name_instruction.pack(
    pady=10
)

patient_name_entry = tk.Entry(
    name_screen,
    width=35,
    font=("Arial", 12)
)

patient_name_entry.pack(
    pady=10
)

patient_name_entry.bind(
    "<KeyRelease>",
    update_dropdown
)

suggestion_list = tk.Listbox(
    name_screen,
    width=35,
    height=5,
    font=("Arial", 11)
)

suggestion_list.bind(
    "<ButtonRelease-1>",
    select_name
)

suggestion_list.bind(
    "<Return>",
    select_name
)

continue_button = tk.Button(
    name_screen,
    text="Continue",
    command=show_data_screen,
    width=15
)

continue_button.pack(
    pady=20
)

name_screen.pack(
    fill="both",
    expand=True
)


# ==========================
# Data Screen
# ==========================

data_screen = tk.Frame(
    window
)

patient_name_label = tk.Label(
    data_screen,
    text="Patient:",
    font=("Arial", 16, "bold")
)

patient_name_label.pack(
    pady=20
)

data_title = tk.Label(
    data_screen,
    text="Patient Measurements",
    font=("Arial", 18)
)

data_title.pack(
    pady=10
)


# ==========================
# Date
# ==========================

date_frame = tk.Frame(
    data_screen
)

date_frame.pack(
    pady=10
)

date_label = tk.Label(
    date_frame,
    text="Date:",
    width=15,
    anchor="e"
)

date_label.pack(
    side="left",
    padx=5
)

date_entry = tk.Entry(
    date_frame,
    width=25
)

date_entry.pack(
    side="left",
    padx=5
)

date_entry.insert(
    0,
    datetime.now().strftime(
        "%d/%m/%Y"
    )
)


# ==========================
# Blood Pressure
# ==========================

bp_frame = tk.Frame(
    data_screen
)

bp_frame.pack(
    pady=10
)

bp_label = tk.Label(
    bp_frame,
    text="Blood Pressure:",
    width=15,
    anchor="e"
)

bp_label.pack(
    side="left",
    padx=5
)


# ==========================
# Systolic
# ==========================

systolic_frame = tk.Frame(
    bp_frame
)

systolic_frame.pack(
    side="left",
    padx=5
)

systolic_entry = tk.Entry(
    systolic_frame,
    width=6
)

systolic_entry.pack()

systolic_label = tk.Label(
    systolic_frame,
    text="Systolic"
)

systolic_label.pack()


# ==========================
# Slash
# ==========================

slash_label = tk.Label(
    bp_frame,
    text="/",
    font=("Arial", 12, "bold")
)

slash_label.pack(
    side="left",
    padx=2
)


# ==========================
# Diastolic
# ==========================

diastolic_frame = tk.Frame(
    bp_frame
)

diastolic_frame.pack(
    side="left",
    padx=5
)

diastolic_entry = tk.Entry(
    diastolic_frame,
    width=6
)

diastolic_entry.pack()

diastolic_label = tk.Label(
    diastolic_frame,
    text="Diastolic"
)

diastolic_label.pack()


# ==========================
# BP Unit
# ==========================

bp_unit_label = tk.Label(
    bp_frame,
    text="mmHg"
)

bp_unit_label.pack(
    side="left",
    padx=5
)


# ==========================
# Heart Rate
# ==========================

hr_frame = tk.Frame(
    data_screen
)

hr_frame.pack(
    pady=10
)

hr_label = tk.Label(
    hr_frame,
    text="Heart Rate:",
    width=15,
    anchor="e"
)

hr_label.pack(
    side="left",
    padx=5
)

hr_entry = tk.Entry(
    hr_frame,
    width=25
)

hr_entry.pack(
    side="left",
    padx=5
)

hr_unit_label = tk.Label(
    hr_frame,
    text="BPM"
)

hr_unit_label.pack(
    side="left",
    padx=5
)


# ==========================
# Risk Warning
# ==========================

risk_warning = tk.Label(
    data_screen,
    text="WARNING\n\n"
         "Your patient is at risk.\n"
         "Please contact your doctor.",
    bg="red",
    fg="white",
    font=("Arial", 12, "bold"),
    padx=20,
    pady=12
)


# ==========================
# Live Risk Checking
# ==========================

systolic_entry.bind(
    "<KeyRelease>",
    check_risk
)

diastolic_entry.bind(
    "<KeyRelease>",
    check_risk
)

hr_entry.bind(
    "<KeyRelease>",
    check_risk
)


# ==========================
# Buttons
# ==========================

button_frame = tk.Frame(
    data_screen
)

button_frame.pack(
    pady=20
)

save_button = tk.Button(
    button_frame,
    text="Data Collect",
    command=save_data,
    width=15
)

save_button.pack(
    side="left",
    padx=10
)

back_button = tk.Button(
    button_frame,
    text="Change Patient",
    command=return_to_name_screen,
    width=15
)

back_button.pack(
    side="left",
    padx=10
)


# ==========================
# Start Program
# ==========================

window.mainloop()