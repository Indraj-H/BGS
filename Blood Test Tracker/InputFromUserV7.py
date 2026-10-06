from openpyxl import Workbook
import tkinter as tk

window = tk.Tk()
window.title("Indy's Data Collector")
window.geometry("700x500")
window.iconbitmap("Icon.ico")

# ==========================
# Title
# ==========================

title = tk.Label(
    window,
    text="Health-E",
    font=("Arial", 16)
)
title.pack()

subtitle = tk.Label(
    window,
    text="Your Health Monitoring Companion",
    font=("Arial", 8),
    fg="gray"
)
subtitle.pack()

# ==========================
# Scrollable Area
# ==========================

container = tk.Frame(window)
container.pack(fill="both", expand=True)

canvas = tk.Canvas(container)
scrollbar = tk.Scrollbar(
    container,
    orient="vertical",
    command=canvas.yview
)

scrollable_frame = tk.Frame(canvas)

canvas_window = canvas.create_window(
    (0, 0),
    window=scrollable_frame,
    anchor="n"
)


def update_scroll_region(event):
    canvas.configure(
        scrollregion=canvas.bbox("all")
    )


def resize_frame(event):
    canvas.itemconfig(
        canvas_window,
        width=event.width
    )


scrollable_frame.bind(
    "<Configure>",
    update_scroll_region
)

canvas.bind(
    "<Configure>",
    resize_frame
)

canvas.configure(
    yscrollcommand=scrollbar.set
)

canvas.pack(
    side="left",
    fill="both",
    expand=True
)

scrollbar.pack(
    side="right",
    fill="y"
)

# ==========================
# People
# ==========================

entries = []


def add_person():
    person_number = len(entries) + 1

    person_frame = tk.Frame(scrollable_frame)
    person_frame.pack(pady=8)

    # Person number above the fields
    label = tk.Label(
        person_frame,
        text=f"Person {person_number}:",
        font=("Arial", 10, "bold")
    )
    label.pack(pady=(0, 5))

    fields_frame = tk.Frame(person_frame)
    fields_frame.pack()

    datelabel = tk.Label(
        fields_frame,
        text="Date(DD/MM/YYYY)"
    )
    datelabel.pack(side="left", padx=5)

    date_entry = tk.Entry(
        fields_frame,
        width=10
    )
    date_entry.pack(side="left", padx=2)

    # Name
    namelabel = tk.Label(
        fields_frame,
        text="Name"
    )
    namelabel.pack(side="left", padx=5)

    name_entry = tk.Entry(
        fields_frame,
        width=20
    )
    name_entry.pack(side="left", padx=5)

    # Blood Pressure
    bplabel = tk.Label(
        fields_frame,
        text="BP(S/D)"
    )
    bplabel.pack(side="left", padx=5)

    systolic_entry = tk.Entry(
        fields_frame,
        width=5
    )
    systolic_entry.pack(side="left", padx=2)

    slashlabel = tk.Label(
        fields_frame,
        text="/"
    )
    slashlabel.pack(side="left", padx=2)

    diastolic_entry = tk.Entry(
        fields_frame,
        width=5
    )
    diastolic_entry.pack(side="left", padx=2)

    hrlabel = tk.Label(
        fields_frame,
        text="HR"
    )
    hrlabel.pack(side="left", padx=2)

    hr_entry = tk.Entry(
        fields_frame,
        width=5
    )
    hr_entry.pack(side="left", padx=2)

    entries.append(
        (
            date_entry,
            name_entry,
            systolic_entry,
            diastolic_entry,
            hr_entry
        )
    )


# ==========================
# Collect Data
# ==========================

def collect_data():
    workbook = Workbook()
    sheet = workbook.active

    sheet.title = "People"

    # Excel headings
    sheet.cell(row=1, column=1).value = "Date(DD/MM/YYYY)"
    sheet.cell(row=1, column=2).value = "Name"
    sheet.cell(row=1, column=3).value = "Systolic BP"
    sheet.cell(row=1, column=4).value = "Diastolic BP"
    sheet.cell(row=1, column=5).value = "Heart Rate"

    # Add people to Excel
    for row, person in enumerate(entries, start=2):
        date_entry = person[0]
        name_entry = person[1]
        systolic_entry = person[2]
        diastolic_entry = person[3]
        hr_entry = person[4]

        sheet.cell(
            row=row,
            column=1
        ).value = date_entry.get()

        sheet.cell(
            row=row,
            column=2
        ).value = name_entry.get()

        sheet.cell(
            row=row,
            column=3
        ).value = systolic_entry.get()

        sheet.cell(
            row=row,
            column=4
        ).value = diastolic_entry.get()

        sheet.cell(
            row=row,
            column=5
        ).value = hr_entry.get()

    workbook.save("names.xlsx")

    success.config(
        text="Data Collected Successfully!",
        fg="green"
    )


# ==========================
# Buttons
# ==========================

button_frame = tk.Frame(window)
button_frame.pack(pady=10)

add_button = tk.Button(
    button_frame,
    text="Add Person",
    command=add_person
)
add_button.pack(side="left", padx=5)

collect_button = tk.Button(
    button_frame,
    text="Collect Data",
    command=collect_data
)
collect_button.pack(side="left", padx=5)

success = tk.Label(
    window,
    text=""
)
success.pack(pady=5)

# ==========================
# Start Program
# ==========================

window.mainloop()
