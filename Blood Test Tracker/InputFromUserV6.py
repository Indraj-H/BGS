from openpyxl import Workbook
import tkinter as tk

window = tk.Tk()
window.title("Indy's Data Collector")
window.geometry("500x500")
window.iconbitmap("Icon.ico")

title = tk.Label(
    window,
    text="Welcome to Indy's Data Collector",
    font=("Arial", 16)
)
title.pack(pady=10)

# ==========================
# Scrollable area
# ==========================

container = tk.Frame(window)
container.pack(fill="both", expand=True)

canvas = tk.Canvas(container)
scrollbar = tk.Scrollbar(
    container,
    orient="vertical",
    command=canvas.yview
)

scrollable_frame = tk.Frame(canvas)

canvas_window = canvas.create_window(
    (0, 0),
    window=scrollable_frame,
    anchor="n"
)

def update_scroll_region(event):
    canvas.configure(scrollregion=canvas.bbox("all"))

def resize_frame(event):
    canvas.itemconfig(
        canvas_window,
        width=event.width
    )

scrollable_frame.bind(
    "<Configure>",
    update_scroll_region
)

canvas.bind(
    "<Configure>",
    resize_frame
)

canvas.configure(
    yscrollcommand=scrollbar.set
)

canvas.pack(
    side="left",
    fill="both",
    expand=True
)

scrollbar.pack(
    side="right",
    fill="y"
)

# ==========================
# People
# ==========================

entries = []


def add_person():
    person_number = len(entries) + 1

    label = tk.Label(
        scrollable_frame,
        text=f"Person {person_number}:"
    )
    label.pack(pady=(5, 2))

    entry = tk.Entry(
        scrollable_frame,
        width=30
    )
    entry.pack(pady=(0, 5))

    entries.append(entry)


# ==========================
# Buttons
# ==========================

button_frame = tk.Frame(window)
button_frame.pack(pady=10)

add_button = tk.Button(
    button_frame,
    text="Add Person",
    command=add_person
)
add_button.pack(side="left", padx=5)


def collect_data():
    workbook = Workbook()
    sheet = workbook.active

    sheet.cell(row=1, column=1).value = "Name"

    for row, entry in enumerate(entries, start=2):
        sheet.cell(row=row, column=1).value = entry.get()

    workbook.save("names.xlsx")

    success.config(
        text="Data Collected Successfully!",
        fg="green"
    )


collect_button = tk.Button(
    button_frame,
    text="Collect Data",
    command=collect_data
)
collect_button.pack(side="left", padx=5)

success = tk.Label(
    window,
    text=""
)
success.pack(pady=5)

window.mainloop()