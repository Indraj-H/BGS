from openpyxl import Workbook, load_workbook
from datetime import datetime
import tkinter as tk
from tkinter import messagebox
import os
import re


window = tk.Tk()
window.title("Indy's Data Collector")
window.geometry("700x550")
window.iconbitmap("Icon.ico")

patient_name = ""
file_name = ""

# ==========================
# File Location
# ==========================

FILE_PATH = r"C:\Users\indra\OneDrive\HSH\Health-E"

os.makedirs(FILE_PATH, exist_ok=True)


# ==========================
# Patient Names
# ==========================

def get_existing_names():
    names = []

    for file in os.listdir(FILE_PATH):
        if file.endswith(".xlsx"):
            name = os.path.splitext(file)[0]
            names.append(name)

    return sorted(names)


existing_names = get_existing_names()

def create_file_name(name):
    safe_name = re.sub(r'[\\/*?:"<>|]', '', name)
    safe_name = safe_name.strip()

    if not safe_name:
        safe_name = "Unknown Patient"

    return f"{safe_name}.xlsx"


def update_dropdown(event=None):
    typed_name = patient_name_entry.get().lower()

    matching_names = [
        name for name in existing_names
        if typed_name in name.lower()
    ]

    suggestion_list.delete(0, tk.END)

    if typed_name and matching_names:
        for name in matching_names:
            suggestion_list.insert(tk.END, name)

        suggestion_list.pack(pady=5)
    else:
        suggestion_list.pack_forget()


def select_name(event=None):
    selected = suggestion_list.curselection()

    if selected:
        selected_name = suggestion_list.get(selected[0])

        patient_name_entry.delete(0, tk.END)
        patient_name_entry.insert(0, selected_name)

        suggestion_list.pack_forget()


# ==========================
# Create Patient Workbook
# ==========================

def create_patient_workbook():

    full_file_path = os.path.join(
        FILE_PATH,
        file_name
    )

    if os.path.exists(full_file_path):
        return

    workbook = Workbook()

    blood_pressure_sheet = workbook.active
    blood_pressure_sheet.title = "Blood Pressure"

    blood_pressure_sheet.cell(
        row=1,
        column=1
    ).value = "Date"

    blood_pressure_sheet.cell(
        row=1,
        column=2
    ).value = "Systolic BP"

    blood_pressure_sheet.cell(
        row=1,
        column=3
    ).value = "Diastolic BP"

    blood_pressure_sheet.cell(
        row=1,
        column=4
    ).value = "Heart Rate"

    blood_test_sheet = workbook.create_sheet(
        "Blood Test"
    )

    blood_test_sheet.cell(
        row=1,
        column=1
    ).value = "Date"

    blood_test_sheet.cell(
        row=1,
        column=2
    ).value = "Blood Test"

    blood_test_sheet.cell(
        row=1,
        column=3
    ).value = "Value"

    workbook.save(
        full_file_path
    )

    workbook.close()


# ==========================
# Risk Warning
# ==========================

def check_risk(event=None):

    systolic = systolic_entry.get().strip()
    diastolic = diastolic_entry.get().strip()
    heart_rate = hr_entry.get().strip()

    if not systolic or not diastolic or not heart_rate:
        risk_warning.pack_forget()
        return

    try:
        systolic_value = int(systolic)
        diastolic_value = int(diastolic)
        heart_rate_value = int(heart_rate)

    except ValueError:
        risk_warning.pack_forget()
        return

    patient_at_risk = (
        systolic_value > 150
        or systolic_value < 100
        or diastolic_value > 90
        or diastolic_value < 70
        or heart_rate_value > 120
        or heart_rate_value < 50
    )

    if patient_at_risk:
        risk_warning.pack(
            pady=10,
            padx=20,
            fill="x"
        )
    else:
        risk_warning.pack_forget()

# ==========================
# Open Patient Excel File
# ==========================

def open_patient_excel():

    full_file_path = os.path.join(
        FILE_PATH,
        file_name
    )

    if os.path.exists(full_file_path):

        os.startfile(full_file_path)

    else:

        messagebox.showerror(
            "File Not Found",
            "The patient's Excel file could not be found."
        )


# ==========================
# Open Excel Button
# ==========================

open_excel_button = tk.Button(
    options_screen,
    text="Open Patient Excel",
    command=open_patient_excel,
    width=20,
    height=2
)

open_excel_button.pack(
    pady=10
)

# ==========================
# Save Blood Pressure
# ==========================

def save_blood_pressure():

    date = bp_date_entry.get().strip()
    systolic = systolic_entry.get().strip()
    diastolic = diastolic_entry.get().strip()
    heart_rate = hr_entry.get().strip()

    if not date or not systolic or not diastolic or not heart_rate:

        messagebox.showerror(
            "Missing Information",
            "Please fill in all fields."
        )

        return

    try:
        datetime.strptime(
            date,
            "%d/%m/%Y"
        )

    except ValueError:

        messagebox.showerror(
            "Invalid Date",
            "Please enter the date in DD/MM/YYYY format."
        )

        return

    try:
        systolic_value = int(systolic)
        diastolic_value = int(diastolic)
        heart_rate_value = int(heart_rate)

    except ValueError:

        messagebox.showerror(
            "Invalid Information",
            "Blood pressure and heart rate must be numbers."
        )

        return

    full_file_path = os.path.join(
        FILE_PATH,
        file_name
    )

    try:

        workbook = load_workbook(
            full_file_path
        )

        sheet = workbook["Blood Pressure"]

        next_row = sheet.max_row + 1

        sheet.cell(
            row=next_row,
            column=1
        ).value = date

        sheet.cell(
            row=next_row,
            column=2
        ).value = systolic_value

        sheet.cell(
            row=next_row,
            column=3
        ).value = diastolic_value

        sheet.cell(
            row=next_row,
            column=4
        ).value = heart_rate_value

        data = list(
            sheet.iter_rows(
                min_row=2,
                values_only=True
            )
        )

        data.sort(
            key=lambda row: datetime.strptime(
                str(row[0]),
                "%d/%m/%Y"
            )
        )

        for row_number, row_data in enumerate(
            data,
            start=2
        ):

            for column_number, value in enumerate(
                row_data,
                start=1
            ):

                sheet.cell(
                    row=row_number,
                    column=column_number
                ).value = value

        workbook.save(
            full_file_path
        )

        workbook.close()

    except PermissionError:

        try:
            workbook.close()
        except:
            pass

        messagebox.showerror(
            "File Error",
            "The Excel file is currently open.\n\n"
            "Please close the file in Excel and try again."
        )

        return

    messagebox.showinfo(
        "Success",
        "Blood pressure data saved!"
    )

    bp_date_entry.delete(0, tk.END)
    systolic_entry.delete(0, tk.END)
    diastolic_entry.delete(0, tk.END)
    hr_entry.delete(0, tk.END)

    bp_date_entry.insert(
        0,
        datetime.now().strftime("%d/%m/%Y")
    )

    risk_warning.pack_forget()


# ==========================
# Save Blood Test
# ==========================

def save_blood_test():

    test_name = blood_test_name_entry.get().strip()
    date = blood_test_date_entry.get().strip()
    value = blood_test_value_entry.get().strip()

    if not test_name or not date or not value:

        messagebox.showerror(
            "Missing Information",
            "Please fill in all fields."
        )

        return

    try:

        datetime.strptime(
            date,
            "%d/%m/%Y"
        )

    except ValueError:

        messagebox.showerror(
            "Invalid Date",
            "Please enter the date in DD/MM/YYYY format."
        )

        return

    full_file_path = os.path.join(
        FILE_PATH,
        file_name
    )

    try:

        workbook = load_workbook(
            full_file_path
        )

        sheet = workbook["Blood Test"]

        # Make sure older files use the new column order
        sheet.cell(
            row=1,
            column=1
        ).value = "Date"

        sheet.cell(
            row=1,
            column=2
        ).value = "Blood Test"

        sheet.cell(
            row=1,
            column=3
        ).value = "Value"

        next_row = sheet.max_row + 1

        sheet.cell(
            row=next_row,
            column=1
        ).value = date

        sheet.cell(
            row=next_row,
            column=2
        ).value = test_name

        sheet.cell(
            row=next_row,
            column=3
        ).value = value

        data = list(
            sheet.iter_rows(
                min_row=2,
                values_only=True
            )
        )

        data.sort(
            key=lambda row: datetime.strptime(
                str(row[0]),
                "%d/%m/%Y"
            )
        )

        for row_number, row_data in enumerate(
            data,
            start=2
        ):

            for column_number, value_data in enumerate(
                row_data,
                start=1
            ):

                sheet.cell(
                    row=row_number,
                    column=column_number
                ).value = value_data

        workbook.save(
            full_file_path
        )

        workbook.close()

    except PermissionError:

        try:
            workbook.close()
        except:
            pass

        messagebox.showerror(
            "File Error",
            "The Excel file is currently open.\n\n"
            "Please close the file in Excel and try again."
        )

        return

    messagebox.showinfo(
        "Success",
        "Blood test data saved!"
    )

    blood_test_name_entry.delete(
        0,
        tk.END
    )

    blood_test_date_entry.delete(
        0,
        tk.END
    )

    blood_test_value_entry.delete(
        0,
        tk.END
    )

    blood_test_date_entry.insert(
        0,
        datetime.now().strftime("%d/%m/%Y")
    )


# ==========================
# Screen Switching
# ==========================

def show_data_options():

    global patient_name
    global file_name
    global existing_names

    patient_name = patient_name_entry.get().strip()

    if not patient_name:

        messagebox.showerror(
            "Missing Name",
            "Please enter the patient's name."
        )

        return

    file_name = create_file_name(
        patient_name
    )

    if patient_name not in existing_names:

        existing_names.append(
            patient_name
        )

        existing_names.sort()

    create_patient_workbook()

    patient_name_options_label.config(
        text=f"Patient: {patient_name}"
    )

    name_screen.pack_forget()

    options_screen.pack(
        fill="both",
        expand=True
    )


def show_blood_pressure():

    options_screen.pack_forget()

    bp_screen.pack(
        fill="both",
        expand=True
    )


def show_blood_test():

    options_screen.pack_forget()

    blood_test_screen.pack(
        fill="both",
        expand=True
    )


def return_to_options():

    bp_screen.pack_forget()
    blood_test_screen.pack_forget()

    risk_warning.pack_forget()

    options_screen.pack(
        fill="both",
        expand=True
    )


def return_to_name_screen():

    options_screen.pack_forget()
    bp_screen.pack_forget()
    blood_test_screen.pack_forget()

    risk_warning.pack_forget()

    patient_name_entry.delete(
        0,
        tk.END
    )

    suggestion_list.pack_forget()

    name_screen.pack(
        fill="both",
        expand=True
    )


# ==========================
# Name Screen
# ==========================

name_screen = tk.Frame(
    window
)

name_title = tk.Label(
    name_screen,
    text="Health-E",
    font=("Arial", 20, "bold")
)

name_title.pack(
    pady=30
)

name_instruction = tk.Label(
    name_screen,
    text="Please enter the patient's name:"
)

name_instruction.pack(
    pady=10
)

patient_name_entry = tk.Entry(
    name_screen,
    width=35,
    font=("Arial", 12)
)

patient_name_entry.pack(
    pady=10
)

patient_name_entry.bind(
    "<KeyRelease>",
    update_dropdown
)

suggestion_list = tk.Listbox(
    name_screen,
    width=35,
    height=5,
    font=("Arial", 11)
)

suggestion_list.bind(
    "<ButtonRelease-1>",
    select_name
)

suggestion_list.bind(
    "<Return>",
    select_name
)

continue_button = tk.Button(
    name_screen,
    text="Continue",
    command=show_data_options,
    width=15
)

continue_button.pack(
    pady=20
)

name_screen.pack(
    fill="both",
    expand=True
)


# ==========================
# Options Screen
# ==========================

options_screen = tk.Frame(
    window
)

patient_name_options_label = tk.Label(
    options_screen,
    text="Patient:",
    font=("Arial", 16, "bold")
)

patient_name_options_label.pack(
    pady=40
)

options_title = tk.Label(
    options_screen,
    text="What would you like to collect?",
    font=("Arial", 18)
)

options_title.pack(
    pady=20
)

blood_pressure_button = tk.Button(
    options_screen,
    text="Blood Pressure",
    command=show_blood_pressure,
    width=20,
    height=2
)

blood_pressure_button.pack(
    pady=10
)

blood_test_button = tk.Button(
    options_screen,
    text="Blood Test",
    command=show_blood_test,
    width=20,
    height=2
)

blood_test_button.pack(
    pady=10
)

change_patient_button = tk.Button(
    options_screen,
    text="Change Patient",
    command=return_to_name_screen,
    width=20
)

change_patient_button.pack(
    pady=30
)


# ==========================
# Blood Pressure Screen
# ==========================

bp_screen = tk.Frame(
    window
)

bp_title = tk.Label(
    bp_screen,
    text="Blood Pressure",
    font=("Arial", 18, "bold")
)

bp_title.pack(
    pady=15
)


# ==========================
# BP Date
# ==========================

bp_date_frame = tk.Frame(
    bp_screen
)

bp_date_frame.pack(
    pady=8
)

bp_date_label = tk.Label(
    bp_date_frame,
    text="Date:",
    width=15,
    anchor="e"
)

bp_date_label.pack(
    side="left",
    padx=5
)

bp_date_entry = tk.Entry(
    bp_date_frame,
    width=25
)

bp_date_entry.pack(
    side="left",
    padx=5
)

bp_date_entry.insert(
    0,
    datetime.now().strftime("%d/%m/%Y")
)


# ==========================
# Blood Pressure
# ==========================

bp_frame = tk.Frame(
    bp_screen
)

bp_frame.pack(
    pady=8
)

bp_label = tk.Label(
    bp_frame,
    text="Blood Pressure:",
    width=15,
    anchor="e"
)

bp_label.pack(
    side="left",
    padx=5
)

systolic_frame = tk.Frame(
    bp_frame
)

systolic_frame.pack(
    side="left",
    padx=5
)

systolic_entry = tk.Entry(
    systolic_frame,
    width=6
)

systolic_entry.pack()

systolic_label = tk.Label(
    systolic_frame,
    text="Systolic"
)

systolic_label.pack()

slash_label = tk.Label(
    bp_frame,
    text="/",
    font=("Arial", 12, "bold")
)

slash_label.pack(
    side="left",
    padx=2
)

diastolic_frame = tk.Frame(
    bp_frame
)

diastolic_frame.pack(
    side="left",
    padx=5
)

diastolic_entry = tk.Entry(
    diastolic_frame,
    width=6
)

diastolic_entry.pack()

diastolic_label = tk.Label(
    diastolic_frame,
    text="Diastolic"
)

diastolic_label.pack()

bp_unit_label = tk.Label(
    bp_frame,
    text="mmHg"
)

bp_unit_label.pack(
    side="left",
    padx=5
)


# ==========================
# Heart Rate
# ==========================

hr_frame = tk.Frame(
    bp_screen
)

hr_frame.pack(
    pady=8
)

hr_label = tk.Label(
    hr_frame,
    text="Heart Rate:",
    width=15,
    anchor="e"
)

hr_label.pack(
    side="left",
    padx=5
)

hr_entry = tk.Entry(
    hr_frame,
    width=25
)

hr_entry.pack(
    side="left",
    padx=5
)

hr_unit_label = tk.Label(
    hr_frame,
    text="BPM"
)

hr_unit_label.pack(
    side="left",
    padx=5
)


# ==========================
# Risk Warning
# ==========================

risk_warning = tk.Label(
    bp_screen,
    text="WARNING\n\n"
         "Your patient is at risk.\n"
         "Please contact your doctor.",
    bg="red",
    fg="white",
    font=("Arial", 12, "bold"),
    padx=20,
    pady=12
)

systolic_entry.bind(
    "<KeyRelease>",
    check_risk
)

diastolic_entry.bind(
    "<KeyRelease>",
    check_risk
)

hr_entry.bind(
    "<KeyRelease>",
    check_risk
)


# ==========================
# Blood Pressure Buttons
# ==========================

bp_button_frame = tk.Frame(
    bp_screen
)

bp_button_frame.pack(
    pady=20
)

bp_save_button = tk.Button(
    bp_button_frame,
    text="Data Collect",
    command=save_blood_pressure,
    width=15
)

bp_save_button.pack(
    side="left",
    padx=10
)

bp_back_button = tk.Button(
    bp_button_frame,
    text="Back",
    command=return_to_options,
    width=15
)

bp_back_button.pack(
    side="left",
    padx=10
)


# ==========================
# Blood Test Screen
# ==========================

blood_test_screen = tk.Frame(
    window
)

blood_test_title = tk.Label(
    blood_test_screen,
    text="Blood Test",
    font=("Arial", 18, "bold")
)

blood_test_title.pack(
    pady=20
)


# ==========================
# Blood Test Name
# ==========================

blood_test_name_frame = tk.Frame(
    blood_test_screen
)

blood_test_name_frame.pack(
    pady=10
)

blood_test_name_label = tk.Label(
    blood_test_name_frame,
    text="Name of Blood Test:",
    width=20,
    anchor="e"
)

blood_test_name_label.pack(
    side="left",
    padx=5
)

blood_test_name_entry = tk.Entry(
    blood_test_name_frame,
    width=25
)

blood_test_name_entry.pack(
    side="left",
    padx=5
)


# ==========================
# Blood Test Date
# ==========================

blood_test_date_frame = tk.Frame(
    blood_test_screen
)

blood_test_date_frame.pack(
    pady=10
)

blood_test_date_label = tk.Label(
    blood_test_date_frame,
    text="Date:",
    width=20,
    anchor="e"
)

blood_test_date_label.pack(
    side="left",
    padx=5
)

blood_test_date_entry = tk.Entry(
    blood_test_date_frame,
    width=25
)

blood_test_date_entry.pack(
    side="left",
    padx=5
)

blood_test_date_entry.insert(
    0,
    datetime.now().strftime("%d/%m/%Y")
)


# ==========================
# Blood Test Value
# ==========================

blood_test_value_frame = tk.Frame(
    blood_test_screen
)

blood_test_value_frame.pack(
    pady=10
)

blood_test_value_label = tk.Label(
    blood_test_value_frame,
    text="Value:",
    width=20,
    anchor="e"
)

blood_test_value_label.pack(
    side="left",
    padx=5
)

blood_test_value_entry = tk.Entry(
    blood_test_value_frame,
    width=25
)

blood_test_value_entry.pack(
    side="left",
    padx=5
)


# ==========================
# Blood Test Buttons
# ==========================

blood_test_button_frame = tk.Frame(
    blood_test_screen
)

blood_test_button_frame.pack(
    pady=30
)

blood_test_save_button = tk.Button(
    blood_test_button_frame,
    text="Data Collect",
    command=save_blood_test,
    width=15
)

blood_test_save_button.pack(
    side="left",
    padx=10
)

blood_test_back_button = tk.Button(
    blood_test_button_frame,
    text="Back",
    command=return_to_options,
    width=15
)

blood_test_back_button.pack(
    side="left",
    padx=10
)


# ==========================
# Start Program
# ==========================

window.mainloop()